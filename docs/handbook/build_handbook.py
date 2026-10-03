"""Build the Vietnamese project handbook from reviewed prose and local evidence."""

from __future__ import annotations

import ast
import json
import re
from datetime import datetime
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
ASSETS = HERE / "assets"
OUTPUT = ROOT / "docs/SO_TAY_KIEN_TRUC_AIRFLOW_PYSPARK_VI.docx"
NAVY = "18334D"
TEAL = "007F86"
PALE = "EDF5F6"
GRAY = "526373"
EVIDENCE = json.loads(
    (ROOT / "docs/evidence/official_run_20260929.json").read_text("utf-8")
)


def md_table(headers, rows):
    rows = [[str(v).replace("|", "/").replace("\n", " ") for v in row] for row in rows]
    return "\n".join(
        ["| " + " | ".join(headers) + " |", "|" + "---|" * len(headers)]
        + ["| " + " | ".join(row) + " |" for row in rows]
    )


def diagram(name, width=12, height=7):
    fig, ax = plt.subplots(figsize=(width, height))
    fig.patch.set_facecolor("white")
    ax.set(xlim=(0, 12), ylim=(0, 8))
    ax.axis("off")
    return fig, ax


def box(ax, x, y, w, h, label, color="#edf5f6", fontsize=13):
    patch = FancyBboxPatch(
        (x, y), w, h, boxstyle="round,pad=0.05,rounding_size=0.14",
        facecolor=color, edgecolor="#007f86", linewidth=1.3,
    )
    ax.add_patch(patch)
    ax.text(x + w / 2, y + h / 2, label, ha="center", va="center",
            fontsize=fontsize, color="#18334d", linespacing=1.4)


def arrow(ax, p1, p2, label=None, color="#526373"):
    ax.add_patch(FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=14,
                                linewidth=1.4, color=color))
    if label:
        ax.text((p1[0]+p2[0])/2, (p1[1]+p2[1])/2 + .15, label,
                ha="center", va="bottom", fontsize=11, color=color,
                bbox=dict(facecolor="white", edgecolor="none", pad=1))


def save_fig(fig, name):
    fig.savefig(ASSETS / f"{name}.png", dpi=210, bbox_inches="tight", pad_inches=.12)
    plt.close(fig)


def create_figures():
    ASSETS.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 13})
    fig, ax = diagram("architecture", height=7.3)
    ax.text(.2, 7.65, "WINDOWS HOST  •  DOCKER DESKTOP / LINUX CONTAINERS",
            fontsize=12, weight="bold", color="#18334d")
    box(ax, .2, 6.15, 3.2, 1, "Trình duyệt\nlocalhost:8080", "#fff3dd")
    box(ax, 4.1, 6.15, 3.6, 1, "airflow-apiserver\nUI / API / Execution API")
    box(ax, 8.5, 6.15, 3.2, 1, "postgres:16\nMetadata + named volume")
    arrow(ax, (3.45, 6.65), (4.05, 6.65))
    arrow(ax, (7.75, 6.65), (8.45, 6.65))
    box(ax, .2, 3.9, 3.2, 1.25, "airflow-dag-processor\nParse định nghĩa DAG")
    box(ax, 4.1, 3.9, 3.6, 1.25, "airflow-scheduler\nScheduler + LocalExecutor")
    box(ax, 8.5, 3.9, 3.2, 1.25, "airflow-init\nMigration / tạo user\nKết thúc exit 0")
    arrow(ax, (3.45, 4.5), (4.05, 4.5))
    arrow(ax, (7.75, 4.7), (8.9, 6.05), "metadata")
    arrow(ax, (10.1, 5.2), (10.1, 6.1))
    box(ax, 4.1, 1.95, 3.6, 1.1, "Spark driver / Python\nlocal[2] • heap 3g", "#d7eeee")
    arrow(ax, (5.9, 3.85), (5.9, 3.1), "spark-submit")
    box(ax, .2, .3, 11.5, .95,
        "BIND MOUNTS: dags / scripts / tests / data / logs / pyproject.toml\n"
        "CSV & Parquet ở data volume; XCom chỉ chứa metadata nhỏ", "#f0f2f5", 10)
    arrow(ax, (1.8, 1.3), (1.8, 3.85), "code DAG")
    arrow(ax, (5.9, 1.9), (5.9, 1.3), "đọc / ghi")
    ax.text(8.1, 2.2, "Không có Spark master/worker riêng.\n"
            "Tách trách nhiệm nhưng chung CPU/RAM.", fontsize=9,
            color="#526373", linespacing=1.6)
    save_fig(fig, "architecture")

    fig, ax = diagram("dag", height=6.8)
    specs = [(4.2, 6.9, 3.6, .7, "start_pipeline\nEmptyOperator"),
             (4.2, 5.5, 3.6, .85, "validate_raw_data\nLandingValidationSensor"),
             (4.2, 4.05, 3.6, .85, "submit_pyspark_etl\nSparkSubmitOperator"),
             (.9, 2.35, 4.1, .9, "compute_rfm_metrics\nSparkSubmitOperator"),
             (7, 2.35, 4.1, .9, "detect_anomalies\nSparkSubmitOperator"),
             (4.2, .55, 3.6, .85, "notify_completion\nPythonOperator • all_success")]
    for values in specs:
        box(ax, *values)
    for start, end in [((6,6.85),(6,6.4)),((6,5.45),(6,4.95)),
                       ((5.1,4),(2.95,3.3)),((6.9,4),(9.05,3.3)),
                       ((2.95,2.3),(5.05,1.45)),((9.05,2.3),(6.95,1.45))]:
        arrow(ax,start,end)
    ax.text(.2,.05,"Hai nhánh độc lập; pool retail_spark=1 nên chạy Spark lần lượt.",
            fontsize=9,color="#526373")
    save_fig(fig,"dag")

    fig, ax = diagram("time_contract", height=4.2)
    box(ax,.4,4.5,4.4,1.45,"D = 2011-12-09\nNgày nghiệp vụ / đầu interval")
    box(ax,7.1,4.5,4.4,1.45,"C = 2011-12-10 00:00\nCutoff loại trừ / cuối interval")
    arrow(ax,(4.9,5.2),(7,5.2),"1 ngày")
    box(ax,.4,1.4,7,1.45,"Lịch sử được đọc: 2010-12-01 → hết D\nĐiều kiện chính: InvoiceDate < C", "#d7eeee")
    box(ax,8,1.4,3.5,1.45,"Chạy thực tế 2026\nkhông thay cutoff", "#fff3dd")
    save_fig(fig,"time_contract")

    fig, ax = diagram("lineage", height=7)
    box(ax,.2,6.4,3.2,1,"CSV gốc CP1252\n541.909 dòng")
    box(ax,4.3,6.4,3.4,1,"Landing UTF-8\nDaily CSV + source manifest")
    arrow(ax,(3.45,6.9),(4.25,6.9),"bootstrap")
    box(ax,8.6,6.4,3.1,1,"ETL theo context\nSHA + parse + dedup")
    arrow(ax,(7.75,6.9),(8.55,6.9))
    box(ax,.2,3.9,3.25,1.1,"transactions\nyear/month • 391.057")
    box(ax,4.4,3.9,3.25,1.1,"audit_input\n536.641 dòng")
    box(ax,8.55,3.9,3.15,1.1,"parse_errors\n0 dòng trong run thật", "#fff3dd")
    for target in (1.8,6.0,10.1):
        arrow(ax,(10.1,6.35),(target,5.05))
    box(ax,.2,1.55,3.25,1.3,"RFM\ncustomers: 4.334\ncluster_profile: 2")
    box(ax,4.4,1.55,7.3,1.3,"AUDIT\nrow_assessments: 536.641\norder_assessments: 25.900 • flagged_orders: 133")
    arrow(ax,(1.8,3.85),(1.8,2.9))
    arrow(ax,(6,3.85),(7.1,2.9))
    ax.text(.2,.5,"Mỗi stage ghi manifest sau khi đủ output; completion mới công bố cả tập kết quả.",
            fontsize=10,color="#526373")
    save_fig(fig,"lineage")

    fig, ax = diagram("publication", height=5.4)
    box(ax,.3,5.25,3.3,1.25,"Attempt ghi generation A\nLỗi giữa write", "#fff0e7")
    box(ax,4.35,5.25,3.3,1.25,"Không có stage commit\nKhông dùng cho consumer", "#fff0e7")
    arrow(ax,(3.65,5.9),(4.3,5.9))
    box(ax,.3,2.65,3.3,1.25,"Attempt ghi generation B\nĐủ Parquet + _SUCCESS")
    box(ax,4.35,2.65,3.3,1.25,"Commit stage manifests\nCùng context / ETL hash")
    box(ax,8.5,2.65,3.2,1.25,"Atomic JSON replace\npublished / C.json", "#d7eeee")
    arrow(ax,(3.65,3.3),(4.3,3.3))
    arrow(ax,(7.7,3.3),(8.45,3.3))
    box(ax,8.5,.55,3.2,1.05,"Consumer đọc\nđúng generation B")
    arrow(ax,(10.1,2.6),(10.1,1.65))
    ax.text(.3,.85,"Con trỏ cũ giữ nguyên đến khi bản mới hoàn chỉnh.\n"
            "Đây là commit metadata local, không phải transaction mọi file.",
            fontsize=9.5, color="#526373",linespacing=1.5)
    save_fig(fig,"publication")

    segments=EVIDENCE["metrics"]["rfm"]["segments"]
    fig, ax = plt.subplots(figsize=(10.5,4.4))
    labels=list(segments)
    values=list(segments.values())
    ax.barh(labels,values,color="#007f86",height=.62)
    for i,value in enumerate(values):
        ax.text(value+18,i,f"{value:,}".replace(",","."),va="center",fontsize=11)
    ax.set_xlim(0,max(values)*1.17)
    ax.set_xlabel("Số khách hàng • tổng 4.334")
    ax.spines[["top","right"]].set_visible(False)
    ax.grid(axis="x",alpha=.15)
    fig.tight_layout()
    save_fig(fig,"segments")

    values=[133,19774,5993]
    fig, ax=plt.subplots(figsize=(10.5,4.2))
    labels=["flagged","not_flagged","not_assessable"]
    ax.barh(labels,values,color=["#bf5b3d","#007f86","#8494a2"],height=.6)
    for i,value in enumerate(values):
        ax.text(value+200,i,f"{value:,}".replace(",","."),va="center",fontsize=11)
    ax.set_xlim(0,23000)
    ax.set_xlabel("Số hóa đơn • tổng 25.900")
    ax.spines[["top","right"]].set_visible(False)
    ax.grid(axis="x",alpha=.15)
    fig.tight_layout()
    save_fig(fig,"order_status")

    timeline=[
        ("start_pipeline","06:17:59.145845","06:17:59.145853"),
        ("validate_raw_data","06:18:01.562749","06:18:54.938081"),
        ("submit_pyspark_etl","06:18:56.043678","06:22:08.033663"),
        ("compute_rfm_metrics","06:22:09.133135","06:24:26.650109"),
        ("detect_anomalies","06:24:27.480240","06:27:21.690396"),
        ("notify_completion","06:27:23.944168","06:27:27.945753"),
    ]
    base=datetime.fromisoformat("2026-09-29T06:17:58.986535")
    fig,ax=plt.subplots(figsize=(11,4.6))
    for i,(label,start,end) in enumerate(timeline):
        left=(datetime.fromisoformat("2026-09-29T"+start)-base).total_seconds()/60
        duration=(datetime.fromisoformat("2026-09-29T"+end)-datetime.fromisoformat("2026-09-29T"+start)).total_seconds()/60
        ax.barh(i,max(duration,.025),left=left,height=.55,color="#007f86" if 2<=i<=4 else "#6586a1")
        ax.text(left+max(duration,.025)+.06,i,f"{duration*60:.1f}s",va="center",fontsize=9)
    ax.set_yticks(range(6),[x[0] for x in timeline])
    ax.invert_yaxis()
    ax.set_xlim(0,10.5)
    ax.set_xlabel("Phút kể từ lúc DagRun bắt đầu • 29/09/2026, UTC")
    ax.grid(axis="x",alpha=.2)
    ax.spines[["top","right"]].set_visible(False)
    fig.tight_layout()
    save_fig(fig,"timeline")
    (ASSETS/"task_timeline.json").write_text(json.dumps(timeline,indent=2)+"\n",encoding="utf-8")


def auto_content():
    content={}
    content["dependencies"]=md_table(
        ["Thành phần", "Phiên bản / nguồn", "Vai trò"], [
        ["Apache Airflow","3.3.1 / base image","Orchestration runtime và SDK"],
        ["Spark provider","6.3.2","SparkSubmitOperator / Hook"],
        ["Standard provider","1.17.0","Empty/Python operators và FileSensor"],
        ["FAB provider","3.8.0","FAB Auth Manager và quản trị user"],
        ["PySpark","4.2.0","DataFrame, SQL, Window và Spark ML"],
        ["Kaggle CLI","Không pin trong requirements.txt","Tải dataset; không phải task hằng ngày"],
        ["pytest","8.4.2","Test runtime và Spark logic"],
        ["Ruff","0.15.7 / requirements-dev.txt","Lint và định dạng code"],
        ["Java","OpenJDK 17 JRE headless","Chạy JVM Spark trong image Linux"],
        ["PostgreSQL","Image postgres:16","Airflow metadata database"],
        ["Dependency tài liệu","requirements-docs.txt","Tạo Word/hình, đọc PDF; tách runtime"],
    ])
    inventory=[
        ("dags/ecommerce_etl_dag.py","Định nghĩa sáu task, Params, date, concurrency và deadline"),
        ("dags/retail_support/tasks.py","Sensor, ngày nghiệp vụ, completion và callbacks"),
        ("dags/retail_support/__init__.py","Package marker cho helper DAG"),
        ("scripts/retail_contracts.py","Contract file, hash, context, stage và publication"),
        ("scripts/prepare_daily_landing.py","Bootstrap nguồn CSV bất biến theo ngày"),
        ("scripts/retail_pipeline.py","Ba stage Spark, outputs/metrics/manifests"),
        ("scripts/pyspark_clean.py","Cleaning legacy và entry ETL pipeline mode"),
        ("scripts/pyspark_rfm.py","RFM, scores, KMeans và validator"),
        ("scripts/pyspark_anomalies.py","Chín row rules, context và bảo toàn dòng"),
        ("scripts/select_kmeans_k.py","Chọn k offline, CSV và chart tùy dependency"),
        ("notebooks/EDA_Online_Retail.ipynb","Khảo sát nguồn/cleaning, không được DAG chạy"),
        ("notebooks/EDA_Anomalies.ipynb","Nghiên cứu missing ID và quy tắc hồi cứu"),
        ("notebooks/EDA_Anomalies_Guide_VI.md","Giải thích các bảng/biểu đồ EDA"),
        ("tests/ và tests/fixtures/","Test logic, DAG, contracts và dữ liệu tổng hợp nhỏ"),
        ("Dockerfile; docker-compose.yaml","Image và topology dịch vụ/volume/network"),
        ("requirements.txt; requirements-dev.txt","Dependency chạy pipeline và lint"),
        ("pyproject.toml","Ruff target Python/line length/rules/ngoại lệ legacy"),
        (".env; .env.example","Environment local / mẫu; không đưa secret vào báo cáo"),
        ("data/raw/","CSV gốc, landing và manifest nguồn"),
        ("data/curated/","Clean snapshots và audit input trung gian"),
        ("data/analytics/rfm_daily/","Customer metrics và profile theo cutoff/version"),
        ("data/analytics/k_selection/","Kết quả nghiên cứu số cụm offline"),
        ("data/audit/anomalies/","Audit chính thức versioned"),
        ("data/audit/anomalies.parquet/","Audit input legacy; khác root audit chính thức"),
        ("data/audit/anomaly_results*/","Kết quả detector legacy và verification cũ"),
        ("data/manifests/","Run contexts, stages, completion, published và locks"),
        ("data/verification/","Artifacts kiểm chứng local, không phải output consumer"),
        ("logs/","Log Airflow/Spark; không đưa toàn bộ vào Git"),
        ("docs/REPORT.md; SETUP.md","Điều hướng báo cáo và hướng dẫn môi trường"),
        ("docs/DAG_DESIGN_VI.md","Thiết kế có nguồn sách, traceability và acceptance plan"),
        ("docs/IMPLEMENTATION_VI.md; evidence/","Hiện trạng và số liệu thực tế"),
        ("docs/handbook/; docs/images/","Nguồn Word, hình minh họa và chỗ lưu hình"),
        ("README.md; CONTRIBUTING.md; AGENTS.md","Giới thiệu và quy tắc cộng tác"),
        (".agent/","TODO, HANDOFF, DECISIONS, plans và quy ước trạng thái"),
        (".github/PULL_REQUEST_TEMPLATE.md","Mẫu nội dung review PR"),
        (".codex/; .agents/skills/","Cấu hình/placeholder hỗ trợ cộng tác, không phải runtime"),
        (".gitignore; .gitkeep","Kiểm soát file versioned và giữ thư mục trống"),
    ]
    content["inventory"]=md_table(["Đường dẫn / nhóm","Trách nhiệm"],inventory)
    tests=[]
    for path in sorted((ROOT/"tests").glob("test_*.py")):
        tree=ast.parse(path.read_text("utf-8"))
        names=[n.name for n in ast.walk(tree) if isinstance(n,ast.FunctionDef) and n.name.startswith("test_")]
        tests.append([path.name,str(len(names)), ", ".join(n.removeprefix("test_") for n in names)])
    # List readable names in small paragraphs rather than one giant table cell.
    content["tests_catalog"]="\n\n".join(
        f"#### {name}\n\nCó {count} hàm test khai báo; parametrize có thể làm số test cases thực thi lớn hơn.\n\n"
        + "\n".join("- "+n for n in names.split(", "))
        for name,count,names in tests
    )
    data=[
        ["DagRun","official_smoke_20260929"], ["Business date / cutoff","2011-12-09 / 2011-12-10"],
        ["Source version",EVIDENCE["context"]["source_version"]],
        ["Task success","6/6"],["Tests","87 passed; 1 upstream warning"],
        ["Raw / dedup / clean","541.909 / 536.641 / 391.057"],
        ["Parse errors","0"],["Khách RFM / cụm","4.334 / 2"],
        ["High-value / inactive","1.734 / 1.463"],
        ["Hóa đơn: flagged / not_flagged / not_assessable","133 / 19.774 / 5.993"],
        ["Dòng có cờ / không có cờ","40.051 / 496.590"],
        ["Monthly partitions","13"],["Independent read-back",EVIDENCE["readback"]["readback"]],
        ["Source SHA-256",EVIDENCE["readback"]["source_sha256"]],
        ["Code fingerprint",EVIDENCE["context"]["code_version"]],
    ]
    content["evidence_summary"]=md_table(["Chỉ số / định danh","Giá trị"],data)
    book=[
        ["B1","Ch.1 §1.1–1.3; 4–19","32–47","DAG, phạm vi điều phối"],
        ["B2","Ch.2 §2.2/2.5; 27–29, 38–40","55–57, 66–68","Operators/tasks/failures"],
        ["B3","Ch.3 §3.4–3.7; 55–66","83–94","Interval, backfill, idempotency"],
        ["B4","Ch.4 §4.2–4.8; 70–85","98–113","Asset scheduling, phương án khác"],
        ["B5","Ch.5 §5.2–5.3; 89–101","117–129","Templating và runtime context"],
        ["B6","Ch.6 §6.1/6.4–6.5; 112–114, 128–136","140–142, 156–164","Dependencies, trigger rules, XCom"],
        ["B7","Ch.7 §7.1; 148–156","176–184","Sensors, timeout, reschedule"],
        ["B8","Ch.8 §8.3.3; 185–186","213–214","Offload compute và SparkSubmit"],
        ["B9","Ch.9 §9.2/9.4; 197–204, 210–212","225–232, 238–240","Connections/hooks/custom sensors"],
        ["B10","Ch.10 §10.1/10.4; 227–240, 251–255","255–268, 279–283","Các lớp kiểm thử"],
        ["B11","Ch.11 §11.2–11.5; 264–287","292–315","Containers/volumes/deployment"],
        ["B12","Ch.12 §12.1–12.4; 296–320","324–348","Best practices, storage, pool"],
        ["B15","Ch.15 §15.1–15.8; 383–418","411–446","Executor, concurrency, monitoring"],
        ["B16","Ch.16 §16.1–16.5; 425–440","453–468","Security và credentials"],
    ]
    content["book_references"]=md_table(["Mã","Chương/mục; trang in","Trang PDF","Áp dụng"],book)
    glossary=[
        ["DAG","Đồ thị dependency không chu trình"],["DagRun","Một lần chạy DAG có run_id và state"],
        ["TaskInstance","Task cụ thể trong một run/attempt"],["Operator","Lớp định nghĩa loại công việc"],
        ["Sensor","Operator chờ điều kiện"],["Hook","Client tích hợp hệ thống"],
        ["Connection","Cấu hình kết nối có ID"],["XCom","Giá trị nhỏ trao đổi giữa task instances"],
        ["Params","Tham số DAG có schema và giá trị theo run"],["Executor","Cơ chế thực thi Airflow task"],
        ["Pool","Giới hạn tài nguyên logic qua số slot"],["Data interval","Khoảng dữ liệu mà scheduled run đại diện"],
        ["Logical date","Định danh thời gian logic; không luôn là thời điểm chạy thật"],
        ["Backfill","Tạo run cho khoảng lịch sử"],["Catchup","Chính sách scheduler tạo interval bị bỏ lỡ"],
        ["Idempotency","Rerun/retry không làm sai kết quả consumer do tác dụng phụ lặp"],
        ["Atomic publication","Chuyển con trỏ metadata hoàn chỉnh như một thao tác thay thế"],
        ["Manifest","Metadata về nguồn/output, version, paths và metrics"],
        ["Generation","Một phiên bản output riêng của một lần thử"],
        ["Partition pruning","Bỏ đọc partition không phù hợp điều kiện"],
        ["Shuffle","Phân phối lại dữ liệu giữa partition để group/join/sort"],
        ["Spark driver","Tiến trình quản lý Spark application"],["Spark task","Đơn vị công việc bên trong Spark, khác Airflow task"],
        ["Window","Phép tính trên cửa sổ/sắp hạng các hàng liên quan"],
        ["RFM","Recency, Frequency, Monetary"],["IQR","Khoảng tứ phân vị Q3−Q1"],
        ["Inactivity proxy","Dấu hiệu dựa trên thời gian chưa mua, không phải xác suất churn"],
        ["Not assessable","Không đủ điều kiện kết luận ở cấp tổng"],
        ["Not applied","Một check không áp dụng với record cụ thể"],
        ["Read-back","Đọc lại output thực tế để đối chiếu contract/số liệu"],
        ["Bind mount","Ánh xạ đường dẫn host vào container"],
        ["Named volume","Lưu trữ được Docker quản lý theo tên"],
    ]
    content["glossary"]=md_table(["Thuật ngữ","Ý nghĩa trong bài"],glossary)
    config=[
        ["DAG ID","ecommerce_etl_dag","dags/ecommerce_etl_dag.py"],
        ["Schedule / timezone","0 0 * * * / UTC","CronDataIntervalTimetable"],
        ["start_date / end_date","2010-12-01 / 2011-12-09","Giới hạn lịch sử"],
        ["catchup / depends_on_past","False / False","Snapshot độc lập, tránh tạo hàng loạt"],
        ["max_active_runs / max_active_tasks","1 / 2","Concurrency cấp DAG"],
        ["Spark master / deploy mode","local[2] / client","Compose connection và operator"],
        ["Driver heap / shuffle partitions","3g / 8","spark_job / run_stage"],
        ["Spark pool","retail_spark; 1 slot mỗi task","Pool database đã tạo 1 slot"],
        ["Source Param","online-retail-v1","Token regex 1–80 ký tự"],
        ["rfm_k","Mặc định 2; từ 2 đến 20","Còn cần đủ khách/vector phân biệt"],
        ["churn_days","Mặc định 90; từ 1 đến 3650","Quy tắc inactivity"],
        ["business_date","Manual bắt buộc; YYYY-MM-DD","C = D + 1 ngày"],
        ["KMeans seed","42","pyspark_rfm.py"],
        ["Row IQR","3×IQR; minimum 30","build_output defaults"],
        ["Order threshold","mean + 3×sample std; minimum 30","assess_orders defaults; không phải UI Params"],
        ["Rule/schema versions","retail-v2 / 1","retail_contracts.py"],
        ["Trusted data root","/opt/airflow/data","RETAIL_DATA_ROOT optional deployment env"],
        ["Metadata size guard","2 MiB","read_json"],
        ["Notification","Local logs","Không cấu hình SMTP/Slack"],
    ]
    content["configuration_reference"]=md_table(["Tên","Giá trị","Nơi/quy tắc"],config)

    schemas=json.loads((HERE/"output_schemas.json").read_text("utf-8"))
    descriptions={
        "InvoiceNo":"Mã hóa đơn gốc", "StockCode":"Mã sản phẩm/phí gốc", "Description":"Mô tả gốc",
        "Quantity":"Số lượng đã parse", "InvoiceDate":"Timestamp giao dịch", "UnitPrice":"Đơn giá gốc sau parse",
        "CustomerID":"ID khách dạng string", "Country":"Quốc gia nguồn", "year":"Partition năm", "month":"Partition tháng",
        "run_date":"Cutoff exclusive C", "LastPurchaseDate":"Lần mua gần nhất trước C", "Recency":"Số ngày từ last purchase đến C",
        "Frequency":"Số invoice phân biệt", "Monetary":"Tổng tiền dương theo nhánh clean", "Cluster":"ID KMeans tùy ý",
        "R_score":"Điểm tương đối Recency 1–5", "F_score":"Điểm Frequency 1–5", "M_score":"Điểm Monetary 1–5",
        "RFM_score":"Chuỗi nối R/F/M score", "is_high_value":"M_score >= 4", "is_inactive":"Recency >= churn_days",
        "observed_history_days":"Độ dài lịch sử nguồn đến cutoff", "churn_threshold_days":"Ngưỡng inactivity được chọn",
        "segment_rule_version":"Phiên bản quy tắc nhãn", "churn_assessment_status":"insufficient_history / inactivity_proxy / not_flagged",
        "segment":"Nhãn nghiệp vụ ưu tiên theo quy tắc", "customers":"Số khách trong cụm",
        "reference_mode":"retrospective_before_run_date", "iqr_multiplier":"Hệ số IQR", "min_samples":"Số mẫu sản phẩm tối thiểu",
        "line_value":"Quantity × UnitPrice hữu hạn", "data_quality_flags":"Danh sách vấn đề trường dữ liệu",
        "context_flags":"Các cờ bối cảnh", "check_results":"Chín kết quả check có cấu trúc", "anomaly_flags":"Tên check flagged",
        "assessment_status":"Trạng thái tổng của dòng", "invoice_key":"InvoiceNo trim/upper dùng grouping",
        "line_count":"Số dòng trong invoice", "invalid_lines":"Số dòng làm tổng đơn không đủ điều kiện",
        "customer_ids":"Số ID khách khác nhau không rỗng", "countries":"Số Country khác nhau", "invoice_days":"Số ngày giao dịch khác nhau",
        "eligible":"Invoice đáp ứng điều kiện nghiệp vụ để lấy tổng", "order_total":"Tổng tiền làm tròn; null nếu invoice không hợp lệ",
        "reason":"Lý do không đủ điều kiện", "reference_n":"Số invoice đủ điều kiện trong baseline",
        "reference_mean":"Mean tổng đơn baseline", "reference_stddev":"Sample standard deviation", "upper_bound":"Ngưỡng trên",
        "currency":"GBP", "multiplier":"Hệ số độ lệch chuẩn", "status":"flagged / not_flagged / not_assessable",
        "rule":"Tên kiểm tra", "observed_value":"Giá trị quan sát", "reference_count":"Số mẫu tham chiếu",
        "q1":"Tứ phân vị 25%", "q3":"Tứ phân vị 75%",
    }
    schema_parts=[]
    for stage, outputs in schemas.items():
        for name,schema in outputs.items():
            schema_parts.append(f"### {stage} / {name}")
            rows=[]
            for field in schema["fields"]:
                datatype=field["type"]
                shown=datatype if isinstance(datatype,str) else datatype["type"]
                if isinstance(datatype,dict) and datatype["type"]=="array":
                    element=datatype["elementType"]
                    shown="array<"+(element if isinstance(element,str) else element["type"])+">"
                desc=descriptions.get(field["name"], "")
                if field["name"].startswith("raw_"):
                    desc="Giá trị chuỗi nguồn trước parse: "+field["name"][4:]
                if field["name"].startswith("mean_"):
                    desc="Trung bình "+field["name"][5:]+" trong cụm"
                rows.append([field["name"],shown,"Có" if field["nullable"] else "Không",desc])
            schema_parts.append(md_table(["Cột","Kiểu Parquet/Spark","Nullable","Ý nghĩa"],rows))
            if name=="row_assessments":
                check=next(f for f in schema["fields"] if f["name"]=="check_results")["type"]["elementType"]
                schema_parts.append("Các trường bên trong mỗi phần tử check_results:")
                schema_parts.append(md_table(["Trường con","Kiểu","Ý nghĩa"],[
                    [f["name"],f["type"],("flagged / not_flagged / not_applied" if f["name"]=="status" else descriptions.get(f["name"], "Trạng thái / bằng chứng của check"))]
                    for f in check["fields"]]))
    content["output_schemas"]="\n\n".join(schema_parts)

    files=["dags/ecommerce_etl_dag.py","dags/retail_support/tasks.py",
           "scripts/retail_contracts.py","scripts/prepare_daily_landing.py","scripts/retail_pipeline.py",
           "scripts/pyspark_clean.py","scripts/pyspark_rfm.py","scripts/pyspark_anomalies.py","scripts/select_kmeans_k.py"]
    vn={
        "spark_job":"Tạo operator Spark với cấu hình chung và timeout của stage.",
        "business_date":"Tách quy tắc manual và scheduled/backfill; chặn ngày không nhất quán.",
        "execute":"Chờ manifest bằng lớp cha rồi kiểm tra nguồn và tạo context.",
        "notify_completion":"Xác minh stage manifests, công bố kết quả và log summary.",
        "log_failure":"Ghi dag/task/run/attempt khi task thất bại.",
        "log_retry":"Ghi sự kiện task thử lại.","log_deadline_miss":"Đọc context deadline giới hạn và ghi cảnh báo.",
        "data_root":"Lấy root đáng tin cậy từ môi trường hoặc mặc định.",
        "safe_path":"Resolve và chặn đường dẫn thoát root.","token":"Kiểm tra nhãn source version bằng regex.",
        "iso_date":"Yêu cầu ngày ISO chuẩn.","days":"Sinh các ngày liên tiếp bao gồm hai biên.",
        "digest":"Băm file SHA-256 theo chunk.","json_hash":"Fingerprint JSON sắp khóa.",
        "read_json":"Đọc metadata có guard kích thước.","atomic_json":"Ghi tạm/fsync rồi thay file metadata.",
        "code_version":"Fingerprint năm file business/shared code.",
        "validate_landing":"Kiểm tra version, coverage, paths, bytes, header và tùy chọn full hash.",
        "context_file":"Định vị context từ run_key hex hợp lệ.",
        "build_context":"Cố định source/code/date/params của DagRun; từ chối thay đổi.",
        "load_context":"Đọc context và xác minh fingerprint code hiện tại.",
        "stage_file":"Định vị manifest của ETL/RFM/audit.",
        "load_stage":"Kiểm tra context hash, complete và marker output.",
        "publication_lock":"Lock exclusive theo cutoff, giải phóng ở finally.",
        "publish_completion":"Chỉ công bố ba stage nhất quán với cùng ETL generation.",
        "prepare":"Bootstrap CP1252 → daily UTF-8, round-trip và manifest bất biến.",
        "finite":"Biểu thức nhận số khác null, NaN và infinity.",
        "read_landing":"Đọc CSV rõ schema và giữ raw values của lỗi parse.",
        "customer_labels":"Gán high-value, inactivity và segment theo thứ tự ưu tiên.",
        "assess_orders":"Tổng hóa đơn có eligibility/reason và ngưỡng mean+3s.",
        "counts":"Collect bảng đếm nhóm nhỏ để ghi metrics.",
        "parquet":"Ghi dataset mới và yêu cầu marker _SUCCESS.",
        "run_clean":"Thực thi ETL pipeline mode, kiểm tra nguồn và ghi ba output.",
        "run_rfm":"Thực thi RFM pipeline mode từ ETL manifest.",
        "run_audit":"Thực thi row/order audit và đối soát số lượng.",
        "run_stage":"Tạo generation, SparkSession, dispatch và commit manifest.",
        "pipeline_main":"CLI --pipeline-context cho từng legacy entry point.",
        "parse_args":"Kiểm tra tham số CLI của script.",
        "read_transactions":"Đọc CSV legacy bằng schema giao dịch tường minh.",
        "prepare_datasets":"Dedup và tách clean-RFM / audit với hai null policies.",
        "write_datasets":"Ghi hai thư mục legacy bằng overwrite.",
        "run_cleaning_job":"Ghép đọc, transform và ghi cleaning legacy.",
        "main":"Điểm vào CLI và vòng đời tài nguyên/error handling.",
        "validate_curated_contract":"Kiểm tra schema và giá trị nhánh RFM trước tổng hợp.",
        "select_history":"Lọc lịch sử theo cutoff; chính sách null-date tùy script.",
        "compute_rfm":"Aggregation theo khách: last purchase, distinct orders, money.",
        "add_rfm_scores":"Window percent_rank giữ ties và tạo điểm 1–5.",
        "prepare_kmeans_features":"log1p ba chỉ số R/F/M.",
        "kmeans_feature_pipeline":"VectorAssembler + StandardScaler thống nhất.",
        "add_rfm_segment":"Fit KMeans có kiểm tra khách/vector/cụm occupied.",
        "validate_rfm_output":"Kiểm tra metrics, scores, cutoff và cluster range.",
        "write_rfm":"Ghi RFM legacy theo ngày và kiểm tra đường dẫn.",
        "validate_paths":"Chặn input/output overlap trước ghi legacy.",
        "validate_input_contract":"Kiểm tra tám cột và đúng kiểu giao dịch cho audit.",
        "_finite":"Biểu thức kiểm tra số hữu hạn dùng trong detector.",
        "_normalized":"Chuẩn hóa chuỗi nội bộ, không thay cột gốc.",
        "_flags":"Tạo array tên flags có điều kiện đúng.",
        "_check":"Tạo struct kết quả một check với reason và evidence.",
        "add_base_columns":"Cột nội bộ, line_value và DQ flags.",
        "add_description_context":"Tên tham chiếu cùng stock và keyword context.",
        "add_invoice_checks":"Sáu quy tắc nghiệp vụ từ nhóm invoice.",
        "add_product_iqr_checks":"Ba rule IQR với eligibility/số mẫu/reason.",
        "build_output":"Ghép chín check, flags và trạng thái tổng của dòng.",
        "validate_output":"Bảo toàn multiset, thứ tự rule và status/cutoff.",
        "write_results":"Ghi partition anomaly legacy của một ngày.",
        "_log_results":"Log bảng tổng hợp trạng thái/rule/reason nhỏ.",
        "score_k_range":"Fit/evaluate các k trên cùng không gian feature.",
        "write_report":"Xuất k_scores.csv và biểu đồ nếu có matplotlib.",
    }
    parts=[]
    for name in files:
        tree=ast.parse((ROOT/name).read_text("utf-8"))
        rows=[]
        for node in ast.walk(tree):
            if isinstance(node,ast.FunctionDef):
                rows.append([node.name, str(node.lineno), vn.get(node.name, ast.get_docstring(node) or "Helper; xem code để biết contract chi tiết.")])
        rows.sort(key=lambda row:int(row[1]))
        parts.append("### "+name+"\n\n"+md_table(["Hàm","Dòng nguồn","Trách nhiệm"],rows))
    content["function_catalog"]="\n\n".join(parts)
    for key,path,lang in [("dockerfile_code","Dockerfile","dockerfile"),
                          ("compose_code","docker-compose.yaml","yaml"),
                          ("dag_code","dags/ecommerce_etl_dag.py","python"),
                          ("lint_code","pyproject.toml","toml")]:
        source=(ROOT/path).read_text("utf-8")
        if key=="compose_code":
            source=source.replace("local-only-change-me","TU_TAO_JWT_SECRET")
        content[key]="```"+lang+"\n"+source.rstrip()+"\n```"
    return content


def set_font(run, name="Calibri", size=None, color=None):
    run.font.name=name
    rpr=run._element.get_or_add_rPr()
    fonts=rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts=OxmlElement("w:rFonts")
        rpr.insert(0,fonts)
    for key in ("ascii","hAnsi","eastAsia","cs"):
        fonts.set(qn("w:"+key),name)
    if size:
        run.font.size=Pt(size)
    if color:
        run.font.color.rgb=RGBColor.from_string(color)


def shade(cell, fill):
    shd=OxmlElement("w:shd")
    shd.set(qn("w:fill"),fill)
    cell._tc.get_or_add_tcPr().append(shd)


def field(paragraph,instruction,text="1"):
    start=OxmlElement("w:fldChar"); start.set(qn("w:fldCharType"),"begin")
    code=OxmlElement("w:instrText"); code.set(qn("xml:space"),"preserve"); code.text=instruction
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    paragraph.add_run()._r.append(start)
    paragraph.add_run()._r.append(code)
    paragraph.add_run()._r.append(sep)
    paragraph.add_run(text)
    paragraph.add_run()._r.append(end)


def inline(paragraph,text):
    # Keep code content copyable. Explicit URLs remain clickable in Word.
    for token in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|https?://[^\s]+)",text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run=paragraph.add_run(token[1:-1]); set_font(run,"Consolas",9.0,TEAL)
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold=True
        elif token.startswith(("https://", "http://")):
            relation=paragraph.part.relate_to(
                token, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
            )
            link=OxmlElement("w:hyperlink"); link.set(qn("r:id"),relation)
            run=OxmlElement("w:r"); props=OxmlElement("w:rPr")
            color=OxmlElement("w:color"); color.set(qn("w:val"),TEAL); props.append(color)
            size=OxmlElement("w:sz"); size.set(qn("w:val"),"18"); props.append(size)
            run.append(props)
            node=OxmlElement("w:t"); node.text=token.replace("/","/\u200b")
            run.append(node); link.append(run); paragraph._p.append(link)
        else:
            paragraph.add_run(token)


class Handbook:
    def __init__(self):
        self.doc=Document()
        self.figures=[]
        self.tables=0
        section=self.doc.sections[0]
        section.page_height=Cm(29.7); section.page_width=Cm(21)
        section.top_margin=Cm(2); section.bottom_margin=Cm(1.9)
        section.left_margin=Cm(2.2); section.right_margin=Cm(2.0)
        section.header_distance=Cm(.85); section.footer_distance=Cm(.85)
        section.different_first_page_header_footer=True
        styles=self.doc.styles
        normal=styles["Normal"]
        normal.font.name="Calibri"; normal.font.size=Pt(11)
        normal.paragraph_format.space_after=Pt(6)
        normal.paragraph_format.line_spacing=1.13
        for level,size,color in [(1,19,NAVY),(2,13,TEAL),(3,11.5,NAVY)]:
            st=styles[f"Heading {level}"]
            st.font.name="Calibri"; st.font.size=Pt(size)
            st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
            st.paragraph_format.keep_with_next=True
            st.paragraph_format.space_before=Pt(12 if level>1 else 0)
            st.paragraph_format.space_after=Pt(7)
            if level==1:
                st.paragraph_format.page_break_before=True
        for name in ("TOC 1","TOC 2"):
            if name not in styles:
                styles.add_style(name,1)
            styles[name].font.name="Calibri"
            styles[name].font.size=Pt(10 if name=="TOC 1" else 9.5)
            styles[name].paragraph_format.space_after=Pt(2)
        code=styles.add_style("Source Code",1)
        code.font.name="Consolas"; code.font.size=Pt(8.0)
        code.paragraph_format.space_after=Pt(0)
        code.paragraph_format.line_spacing=1.0
        code.paragraph_format.left_indent=Cm(.15)
        code.paragraph_format.right_indent=Cm(.1)
        code.paragraph_format.widow_control=False
        caption=styles["Caption"]
        caption.font.name="Calibri"; caption.font.size=Pt(9)
        caption.font.color.rgb=RGBColor.from_string(GRAY)
        caption.font.italic=True
        p=section.header.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        set_font(p.add_run("AIRFLOW × PYSPARK  |  RETAIL DATA PLATFORM"),size=8,color=GRAY)
        p=section.footer.paragraphs[0]
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run("Sổ tay project  •  29/09/2026  |  Trang "),size=8,color=GRAY)
        field(p," PAGE ")
        p.add_run(" / ")
        field(p," NUMPAGES ")
        for run in p.runs:
            set_font(run,size=8,color=GRAY)
        self.doc.core_properties.title="Sổ tay kiến trúc và vận hành Airflow – PySpark Retail"
        self.doc.core_properties.subject="Kiến trúc, thiết lập, DAG, ETL, RFM, audit và bằng chứng kiểm chứng"
        self.doc.core_properties.author="Nhóm project Retail Data Engineering"
        self.doc.core_properties.keywords="Airflow, PySpark, Retail, RFM, DAG, kiến trúc, hướng dẫn"
        self.doc.core_properties.comments="Tạo từ code và evidence local; không chứa secret người dùng."

    def cover(self):
        doc=self.doc
        p=doc.add_paragraph()
        p.paragraph_format.space_before=Cm(1.8)
        set_font(p.add_run("SỔ TAY KIẾN TRÚC\nVÀ VẬN HÀNH"),size=29,color=NAVY)
        p=doc.add_paragraph()
        set_font(p.add_run("APACHE AIRFLOW\nPYSPARK RETAIL"),size=24,color=TEAL)
        p.paragraph_format.space_after=Pt(20)
        p=doc.add_paragraph("Từ yêu cầu bài đến hệ thống chạy thực tế")
        set_font(p.runs[0],size=15,color=GRAY)
        doc.add_paragraph(
            "Tài liệu tiếng Việt dành cho người mới, nhóm phát triển và người bảo vệ bài. "
            "Bao gồm cơ sở lý thuyết, kiến trúc triển khai, từng task và thuật toán, "
            "hướng dẫn từ đầu, xử lý sự cố và bằng chứng nghiệm thu."
        )
        doc.add_paragraph()
        self.table(["Phạm vi","Thông tin"],[
            ["Project","Airflow-Apache-Retail-E-Commerce-Dataset"],
            ["Ngày chốt bằng chứng","29/09/2026"],
            ["Runtime","Airflow 3.3.1 • PySpark 4.2.0 • PostgreSQL 16"],
            ["Run kiểm chứng","official_smoke_20260929 • 6/6 task success"],
            ["Kiểm thử","87 tests passed • 8 Parquet outputs read-back"],
            ["Yêu cầu","airflow_subject.docx + rubric cập nhật 25/25/25/15/10"],
        ],count=False)
        p=doc.add_paragraph("Bản chỉnh sửa chính: DOCX  |  Bản xem/in: PDF đi kèm")
        p.paragraph_format.space_before=Pt(20)
        set_font(p.runs[0],size=9,color=GRAY)
        doc.add_page_break()
        p=doc.add_paragraph("MỤC LỤC")
        set_font(p.runs[0],size=21,color=NAVY)
        doc.add_paragraph("Chọn tiêu đề để đi đến phần cần đọc. Trong Word có thể mở Navigation Pane bằng Ctrl+F.")
        p=doc.add_paragraph()
        field(p,' TOC \\o "1-2" \\h \\z \\u ',"Mục lục tự động được cập nhật khi dựng bản Word cuối.")

    def table(self,headers,rows,count=True):
        if count:
            self.tables+=1
        table=self.doc.add_table(rows=1,cols=len(headers))
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        width=16.8
        weights={2:[.33,.67],3:[.25,.29,.46],4:[.20,.23,.14,.43]}.get(len(headers),[1/len(headers)]*len(headers))
        # Common task/matrix tables need reasonably balanced four-column widths.
        if len(headers)==4 and headers[0] in ("Task","Yêu cầu","Thời gian","Mã"):
            weights=[.19,.28,.22,.31]
        if headers==["Cột","Kiểu Parquet/Spark","Nullable","Ý nghĩa"]:
            weights=[.27,.21,.11,.41]
        if headers==["Hàm","Dòng nguồn","Trách nhiệm"]:
            weights=[.32,.1,.58]
        if headers==["Mã","Nguồn","Cách sử dụng"]:
            weights=[.07,.61,.32]
        if headers==["Mã","Chương/mục; trang in","Trang PDF","Áp dụng"]:
            weights=[.07,.41,.18,.34]
        for col,weight in zip(table.columns,weights):
            col.width=Cm(width*weight)
        for i,text in enumerate(headers):
            table.rows[0].cells[i].text=text
        repeat=OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"),"true")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=table.add_row().cells
            for i,text in enumerate(row):
                cells[i].text=""
                inline(cells[i].paragraphs[0],str(text))
        for ridx,row in enumerate(table.rows):
            cant=OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant)
            for i,cell in enumerate(row.cells):
                cell.width=Cm(width*weights[i])
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade(cell,NAVY if ridx==0 else ("F0F5F7" if ridx%2 else "FFFFFF"))
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.keep_with_next=(ridx==0)
                    for run in p.runs:
                        set_font(run,size=9 if len(headers)<4 else 8.7,color="FFFFFF" if ridx==0 else NAVY)
                        if ridx==0:
                            run.bold=True
        self.doc.add_paragraph().paragraph_format.space_after=Pt(2)

    def markdown(self,text):
        lines=text.splitlines(); index=0
        while index<len(lines):
            line=lines[index].strip()
            if not line or line.startswith("# "):
                index+=1; continue
            if line.startswith("```"):
                index+=1
                while index<len(lines) and not lines[index].strip().startswith("```"):
                    p=self.doc.add_paragraph(style="Source Code")
                    run=p.add_run(lines[index].replace("\t","    "))
                    set_font(run,"Consolas",8.0,NAVY)
                    ppr=p._p.get_or_add_pPr()
                    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"F1F4F7"); ppr.append(shd)
                    index+=1
                self.doc.add_paragraph().paragraph_format.space_after=Pt(2)
                index+=1; continue
            if line.startswith("|"):
                headers=[x.strip() for x in line.strip("|").split("|")]
                index+=2; rows=[]
                while index<len(lines) and lines[index].strip().startswith("|"):
                    row=[x.strip() for x in lines[index].strip().strip("|").split("|")]
                    if len(row)!=len(headers):
                        raise ValueError((line,row))
                    rows.append(row); index+=1
                self.table(headers,rows); continue
            match=re.match(r"!\[([^]]+)\]\(([^)]+)\)",line)
            if match:
                p=self.doc.add_paragraph()
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next=True
                shape=p.add_run().add_picture(str(HERE/match[2]),width=Cm(16.2))
                shape._inline.docPr.set("descr",match[1])
                cap=self.doc.add_paragraph(match[1],style="Caption")
                cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
                self.figures.append(match[1])
                index+=1; continue
            if line.startswith("## "):
                self.doc.add_heading(line[3:],level=1); index+=1; continue
            if line.startswith("### "):
                self.doc.add_heading(line[4:],level=2); index+=1; continue
            if line.startswith("#### "):
                self.doc.add_heading(line[5:],level=3); index+=1; continue
            if line.startswith("- "):
                p=self.doc.add_paragraph(style="List Bullet")
                inline(p,line[2:]); index+=1; continue
            paragraph=[line]; index+=1
            while index<len(lines) and lines[index].strip() and not lines[index].lstrip().startswith(("#","|","```","![","- ")):
                paragraph.append(lines[index].strip()); index+=1
            p=self.doc.add_paragraph()
            inline(p," ".join(paragraph))

    def save(self):
        self.doc.save(OUTPUT)
        return {"output":str(OUTPUT),"paragraphs":len(self.doc.paragraphs),
                "tables":len(self.doc.tables),"figures":len(self.figures)}


def main():
    create_figures()
    text=(HERE/"SO_TAY_KIEN_TRUC_VI.md").read_text("utf-8")
    for key,value in auto_content().items():
        marker="{{"+key+"}}"
        assert marker in text,key
        text=text.replace(marker,value)
    unresolved=re.findall(r"\{\{([a-z_]+)\}\}",text)
    assert not unresolved,unresolved
    (HERE/"SO_TAY_KIEN_TRUC_VI.expanded.md").write_text(text,encoding="utf-8")
    doc=Handbook()
    doc.cover()
    doc.markdown(text)
    summary=doc.save()
    summary["word_count_source"]=len(text.split())
    summary["chapters_and_appendices"]=len(re.findall(r"^## ",text,re.M))
    (HERE/"build_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=True,indent=2))


if __name__=="__main__":
    main()
